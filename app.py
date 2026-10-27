from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    jsonify,
    session,
    flash,
    send_from_directory,
    abort,
)
from werkzeug.utils import secure_filename
from werkzeug.security import check_password_hash, generate_password_hash
from functools import wraps

import sqlite3
import os
import json
import re
import tempfile
import markdown
from time import monotonic
from uuid import uuid4
from secrets import token_urlsafe
from markupsafe import Markup

from ddgs import DDGS
from datetime import date, datetime, timedelta
from dotenv import load_dotenv
from groq import Groq

try:
    from pywebpush import webpush, WebPushException
except ImportError:
    webpush = None
    WebPushException = Exception

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None


# =========================================================
# APPLICATION CONFIGURATION
# =========================================================

load_dotenv()

app = Flask(__name__)
app.config["SECRET_KEY"] = os.getenv(
    "FLASK_SECRET_KEY",
    "smart-lecture-ai-development-key-change-before-deployment"
)
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = os.getenv("COOKIE_SECURE", "0") == "1"

api_key = os.getenv("GROQ_API_KEY")
client = Groq(api_key=api_key) if api_key else None

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATABASE = os.getenv(
    "DATABASE_PATH",
    os.path.join(BASE_DIR, "lecture_notes.db")
)
UPLOAD_FOLDER = os.getenv(
    "UPLOAD_FOLDER",
    os.path.join(BASE_DIR, "uploads")
)
PROFILE_PICTURE_FOLDER = os.path.join(UPLOAD_FOLDER, "profile-pictures")

MAX_AUDIO_SIZE = 25 * 1024 * 1024

ALLOWED_EXTENSIONS = {
    "wav", "mp3", "mp4", "mpeg", "mpga",
    "m4a", "ogg", "webm", "flac"
}
ALLOWED_DOCUMENT_EXTENSIONS = {"pdf", "docx", "pptx", "txt", "md"}
ALLOWED_PROFILE_PICTURE_EXTENSIONS = {"jpg", "jpeg", "png", "webp", "gif"}
PROFILE_AVATAR_OPTIONS = [
    ("preset:book", "📚", "Library"),
    ("preset:rocket", "🚀", "Explorer"),
    ("preset:bulb", "💡", "Inventor"),
    ("preset:atom", "⚛️", "Scientist"),
    ("preset:star", "⭐", "Achiever"),
    ("preset:leaf", "🌿", "Calm focus"),
]
PROFILE_AVATAR_VALUES = {option[0] for option in PROFILE_AVATAR_OPTIONS}
AVAILABLE_THEMES = {
    "dark", "light", "ocean", "forest", "sunset", "lavender"
}
PUSH_VAPID_PUBLIC_KEY = os.getenv("VAPID_PUBLIC_KEY", "")
PUSH_VAPID_PRIVATE_KEY = os.getenv("VAPID_PRIVATE_KEY", "")
PUSH_VAPID_CLAIMS_EMAIL = os.getenv(
    "VAPID_CLAIMS_EMAIL", "mailto:admin@example.com"
)
TEACHER_EMAILS = {
    value.strip().lower()
    for value in os.getenv("TEACHER_EMAILS", "").split(",")
    if value.strip()
}

os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(PROFILE_PICTURE_FOLDER, exist_ok=True)

app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER
app.config["MAX_CONTENT_LENGTH"] = MAX_AUDIO_SIZE

AUTH_RATE_LIMIT_WINDOW = 300
AUTH_RATE_LIMIT_MAX_ATTEMPTS = 10
_auth_attempts = {}
CSRF_SESSION_KEY = "_csrf_token"


# =========================================================
# DATABASE HELPERS
# =========================================================

def get_db_connection():
    connection = sqlite3.connect(DATABASE)
    connection.row_factory = sqlite3.Row
    return connection


def get_csrf_token():
    token = session.get(CSRF_SESSION_KEY)
    if not token:
        token = token_urlsafe(32)
        session[CSRF_SESSION_KEY] = token
    return token


def csrf_field():
    return Markup(
        '<input type="hidden" name="csrf_token" value="{}">'.format(
            get_csrf_token()
        )
    )


@app.before_request
def protect_form_requests():
    if request.method != "POST" or request.path.startswith("/api/"):
        return None

    submitted_token = request.form.get("csrf_token", "")
    expected_token = session.get(CSRF_SESSION_KEY, "")
    if not expected_token or not submitted_token or submitted_token != expected_token:
        return (
            render_template(
                "access_denied.html",
                message="This form expired or came from an untrusted source. Refresh the page and try again."
            ),
            400,
        )
    return None


@app.after_request
def add_security_headers(response):
    response.headers.setdefault("X-Content-Type-Options", "nosniff")
    response.headers.setdefault("X-Frame-Options", "SAMEORIGIN")
    response.headers.setdefault("Referrer-Policy", "strict-origin-when-cross-origin")
    response.headers.setdefault(
        "Permissions-Policy",
        "camera=(), geolocation=(), microphone=(self), payment=()"
    )
    response.headers.setdefault(
        "Content-Security-Policy",
        "default-src 'self'; img-src 'self' data:; "
        "style-src 'self' 'unsafe-inline'; "
        "script-src 'self' 'unsafe-inline'; "
        "connect-src 'self'; font-src 'self' data:; "
        "base-uri 'self'; form-action 'self'; frame-ancestors 'self'"
    )
    if request.is_secure:
        response.headers.setdefault(
            "Strict-Transport-Security",
            "max-age=31536000; includeSubDomains"
        )
    return response


def auth_rate_limited(action):
    """Limit repeated login/registration attempts per client address."""
    now = monotonic()
    key = (action, request.remote_addr or "unknown")
    attempts = [
        timestamp for timestamp in _auth_attempts.get(key, [])
        if now - timestamp < AUTH_RATE_LIMIT_WINDOW
    ]
    if len(attempts) >= AUTH_RATE_LIMIT_MAX_ATTEMPTS:
        _auth_attempts[key] = attempts
        return True
    attempts.append(now)
    _auth_attempts[key] = attempts
    return False


def extract_document_text(uploaded_file, extension):
    """Extract readable text from a supported study document."""
    if extension in {"txt", "md"}:
        return uploaded_file.read().decode("utf-8", errors="replace").strip()

    if extension == "pdf":
        if PdfReader is None:
            raise RuntimeError("PDF support is not installed.")
        reader = PdfReader(uploaded_file)
        return "\n\n".join((page.extract_text() or "") for page in reader.pages).strip()

    if extension == "docx":
        if Document is None:
            raise RuntimeError("DOCX support is not installed.")
        document = Document(uploaded_file)
        return "\n".join(
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ).strip()

    if extension == "pptx":
        if Presentation is None:
            raise RuntimeError("PPTX support is not installed.")
        presentation = Presentation(uploaded_file)
        slides = []
        for slide in presentation.slides:
            text = "\n".join(
                shape.text.strip()
                for shape in slide.shapes
                if hasattr(shape, "text") and shape.text.strip()
            )
            if text:
                slides.append(text)
        return "\n\n".join(slides).strip()

    raise ValueError("Unsupported document type.")


def initialize_database():
    connection = get_db_connection()

    try:
        connection.execute("""
            CREATE TABLE IF NOT EXISTS lectures (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER,
                subject TEXT NOT NULL,
                date TEXT NOT NULL,
                title TEXT NOT NULL,
                content TEXT NOT NULL,
                FOREIGN KEY (user_id) REFERENCES users(id)
            )
        """)

        connection.execute("""
            CREATE TABLE IF NOT EXISTS quiz_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER,
                quiz_type TEXT NOT NULL,
                lecture_id INTEGER,
                lecture_title TEXT NOT NULL,
                score REAL NOT NULL,
                max_score REAL NOT NULL,
                percentage REAL NOT NULL,
                attempted_at TEXT NOT NULL,
                FOREIGN KEY (lecture_id) REFERENCES lectures(id)
                ,FOREIGN KEY (user_id) REFERENCES users(id)
            )
        """)

        connection.execute("""
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                username TEXT UNIQUE,
                email TEXT NOT NULL UNIQUE,
                password_hash TEXT NOT NULL,
                created_at TEXT NOT NULL,
                profile_picture TEXT,
                role TEXT NOT NULL DEFAULT 'student'
            )
        """)

        user_columns = {
            row["name"] for row in connection.execute(
                "PRAGMA table_info(users)"
            ).fetchall()
        }
        if "profile_picture" not in user_columns:
            connection.execute(
                "ALTER TABLE users ADD COLUMN profile_picture TEXT"
            )
        if "role" not in user_columns:
            connection.execute(
                "ALTER TABLE users ADD COLUMN role TEXT NOT NULL DEFAULT 'student'"
            )
        if "username" not in user_columns:
            connection.execute(
                "ALTER TABLE users ADD COLUMN username TEXT"
            )
        missing_usernames = connection.execute(
            "SELECT id, name FROM users WHERE username IS NULL OR username = ''"
        ).fetchall()
        for row in missing_usernames:
            base = re.sub(r"[^a-z0-9_]+", "_", row["name"].lower()).strip("_")
            username = (base or "student")[:18]
            connection.execute(
                "UPDATE users SET username = ? WHERE id = ?",
                (f"{username}_{row['id']}", row["id"])
            )
        for table, column, definition in (
            ("lectures", "user_id", "INTEGER"),
            ("quiz_history", "user_id", "INTEGER"),
        ):
            columns = {
                row["name"] for row in connection.execute(
                    f"PRAGMA table_info({table})"
                ).fetchall()
            }
            if column not in columns:
                connection.execute(
                    f"ALTER TABLE {table} ADD COLUMN {column} {definition}"
                )

        connection.execute("""
            CREATE TABLE IF NOT EXISTS user_settings (
                user_id INTEGER PRIMARY KEY,
                language TEXT NOT NULL DEFAULT 'bilingual',
                notifications_enabled INTEGER NOT NULL DEFAULT 1,
                focus_reminders_enabled INTEGER NOT NULL DEFAULT 1,
                theme TEXT NOT NULL DEFAULT 'dark',
                left_handed_mode INTEGER NOT NULL DEFAULT 0,
                contacts_enabled INTEGER NOT NULL DEFAULT 0,
                FOREIGN KEY (user_id) REFERENCES users(id)
            )
        """)
        settings_columns = {
            row["name"] for row in connection.execute(
                "PRAGMA table_info(user_settings)"
            ).fetchall()
        }
        if "theme" not in settings_columns:
            connection.execute(
                "ALTER TABLE user_settings ADD COLUMN theme TEXT NOT NULL DEFAULT 'dark'"
            )
        if "left_handed_mode" not in settings_columns:
            connection.execute(
                "ALTER TABLE user_settings ADD COLUMN left_handed_mode INTEGER NOT NULL DEFAULT 0"
            )
        if "contacts_enabled" not in settings_columns:
            connection.execute(
                "ALTER TABLE user_settings ADD COLUMN contacts_enabled INTEGER NOT NULL DEFAULT 0"
            )

        connection.execute("""
            CREATE TABLE IF NOT EXISTS friendships (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                requester_id INTEGER NOT NULL,
                addressee_id INTEGER NOT NULL,
                status TEXT NOT NULL DEFAULT 'pending',
                created_at TEXT NOT NULL,
                UNIQUE (requester_id, addressee_id),
                CHECK (requester_id != addressee_id),
                FOREIGN KEY (requester_id) REFERENCES users(id),
                FOREIGN KEY (addressee_id) REFERENCES users(id)
            )
        """)

        connection.execute("""
            CREATE TABLE IF NOT EXISTS messages (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                sender_id INTEGER NOT NULL,
                receiver_id INTEGER NOT NULL,
                body TEXT NOT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY (sender_id) REFERENCES users(id),
                FOREIGN KEY (receiver_id) REFERENCES users(id)
            )
        """)

        connection.execute("""
            CREATE TABLE IF NOT EXISTS study_groups (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                description TEXT NOT NULL DEFAULT '',
                owner_id INTEGER NOT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY (owner_id) REFERENCES users(id)
            )
        """)

        connection.execute("""
            CREATE TABLE IF NOT EXISTS study_group_members (
                group_id INTEGER NOT NULL,
                user_id INTEGER NOT NULL,
                joined_at TEXT NOT NULL,
                PRIMARY KEY (group_id, user_id),
                FOREIGN KEY (group_id) REFERENCES study_groups(id),
                FOREIGN KEY (user_id) REFERENCES users(id)
            )
        """)

        connection.execute("""
            CREATE TABLE IF NOT EXISTS study_group_messages (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                group_id INTEGER NOT NULL,
                sender_id INTEGER NOT NULL,
                body TEXT NOT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY (group_id) REFERENCES study_groups(id),
                FOREIGN KEY (sender_id) REFERENCES users(id)
            )
        """)

        connection.execute("""
            CREATE TABLE IF NOT EXISTS notifications (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                title TEXT NOT NULL,
                body TEXT NOT NULL,
                link TEXT,
                kind TEXT NOT NULL DEFAULT 'general',
                is_read INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL,
                FOREIGN KEY (user_id) REFERENCES users(id)
            )
        """)

        connection.execute("""
            CREATE TABLE IF NOT EXISTS exam_plans (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                subject TEXT NOT NULL,
                exam_date TEXT NOT NULL,
                units TEXT NOT NULL DEFAULT '',
                daily_minutes INTEGER NOT NULL DEFAULT 25,
                confidence INTEGER NOT NULL DEFAULT 3,
                created_at TEXT NOT NULL,
                FOREIGN KEY (user_id) REFERENCES users(id)
            )
        """)

        connection.execute("""
            CREATE TABLE IF NOT EXISTS study_tasks (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                plan_id INTEGER,
                title TEXT NOT NULL,
                due_date TEXT NOT NULL,
                completed INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL,
                FOREIGN KEY (user_id) REFERENCES users(id),
                FOREIGN KEY (plan_id) REFERENCES exam_plans(id)
            )
        """)

        connection.execute("""
            CREATE TABLE IF NOT EXISTS classroom_announcements (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                author_id INTEGER,
                classroom_id INTEGER,
                title TEXT NOT NULL,
                body TEXT NOT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY (author_id) REFERENCES users(id),
                FOREIGN KEY (classroom_id) REFERENCES classrooms(id)
            )
        """)
        announcement_columns = {
            row["name"] for row in connection.execute(
                "PRAGMA table_info(classroom_announcements)"
            ).fetchall()
        }
        if "author_id" not in announcement_columns:
            connection.execute(
                "ALTER TABLE classroom_announcements ADD COLUMN author_id INTEGER"
            )
        if "classroom_id" not in announcement_columns:
            connection.execute(
                "ALTER TABLE classroom_announcements ADD COLUMN classroom_id INTEGER"
            )

        connection.execute("""
            CREATE TABLE IF NOT EXISTS classrooms (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                teacher_id INTEGER NOT NULL,
                name TEXT NOT NULL,
                join_code TEXT NOT NULL UNIQUE,
                created_at TEXT NOT NULL,
                FOREIGN KEY (teacher_id) REFERENCES users(id)
            )
        """)

        connection.execute("""
            CREATE TABLE IF NOT EXISTS classroom_members (
                classroom_id INTEGER NOT NULL,
                user_id INTEGER NOT NULL,
                joined_at TEXT NOT NULL,
                PRIMARY KEY (classroom_id, user_id),
                FOREIGN KEY (classroom_id) REFERENCES classrooms(id),
                FOREIGN KEY (user_id) REFERENCES users(id)
            )
        """)

        connection.execute("""
            CREATE TABLE IF NOT EXISTS user_activity (
                user_id INTEGER PRIMARY KEY,
                last_seen_at TEXT NOT NULL,
                total_active_seconds INTEGER NOT NULL DEFAULT 0,
                FOREIGN KEY (user_id) REFERENCES users(id)
            )
        """)

        connection.execute("""
            CREATE TABLE IF NOT EXISTS push_subscriptions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                endpoint TEXT NOT NULL UNIQUE,
                subscription_json TEXT NOT NULL,
                created_at TEXT NOT NULL,
                last_used_at TEXT NOT NULL,
                FOREIGN KEY (user_id) REFERENCES users(id)
            )
        """)

        count = connection.execute(
            "SELECT COUNT(*) FROM lectures"
        ).fetchone()[0]

        if count == 0:
            connection.execute("""
                INSERT INTO lectures (
                    subject, date, title, content
                )
                VALUES (?, ?, ?, ?)
            """, (
                "Introduction to Electronics",
                "17 October 2026",
                "Basic Electronic Components",
                """1. Resistors limit the flow of electric current.

2. Capacitors store electrical energy in an electric field.

3. Diodes allow current to flow primarily in one direction.

4. Transistors can act as switches or amplifiers."""
            ))

        connection.commit()

    finally:
        connection.close()


def get_current_user():
    user_id = session.get("user_id")

    if not user_id:
        return None

    connection = get_db_connection()

    try:
        return connection.execute(
            "SELECT id, name, username, email, profile_picture, role FROM users WHERE id = ?",
            (user_id,)
        ).fetchone()
    finally:
        connection.close()


def get_user_settings(user_id):
    connection = get_db_connection()

    try:
        settings = connection.execute("""
            SELECT language, notifications_enabled,
                   focus_reminders_enabled, theme, left_handed_mode,
                   contacts_enabled
            FROM user_settings
            WHERE user_id = ?
        """, (user_id,)).fetchone()

        if settings is None:
            connection.execute("""
                INSERT INTO user_settings (user_id)
                VALUES (?)
            """, (user_id,))
            connection.commit()
            settings = connection.execute("""
                SELECT language, notifications_enabled,
                       focus_reminders_enabled, theme, left_handed_mode,
                       contacts_enabled
                FROM user_settings
                WHERE user_id = ?
            """, (user_id,)).fetchone()

        return settings
    finally:
        connection.close()


def record_user_activity(user_id, active_seconds=0):
    now = datetime.now().isoformat(timespec="seconds")
    connection = get_db_connection()
    try:
        connection.execute("""
            INSERT INTO user_activity (
                user_id, last_seen_at, total_active_seconds
            )
            VALUES (?, ?, ?)
            ON CONFLICT(user_id) DO UPDATE SET
                last_seen_at = excluded.last_seen_at,
                total_active_seconds = user_activity.total_active_seconds + excluded.total_active_seconds
        """, (user_id, now, max(0, min(int(active_seconds), 300))))
        connection.commit()
    finally:
        connection.close()


def send_push_to_user(user_id, title, body, link="/notifications"):
    if not (
        webpush
        and PUSH_VAPID_PUBLIC_KEY
        and PUSH_VAPID_PRIVATE_KEY
    ):
        return 0

    connection = get_db_connection()
    try:
        subscriptions = connection.execute("""
            SELECT id, endpoint, subscription_json
            FROM push_subscriptions
            WHERE user_id = ?
        """, (user_id,)).fetchall()
        sent = 0
        for subscription in subscriptions:
            try:
                webpush(
                    subscription_info=json.loads(subscription["subscription_json"]),
                    data=json.dumps({
                        "title": title,
                        "body": body,
                        "link": link,
                    }),
                    vapid_private_key=PUSH_VAPID_PRIVATE_KEY,
                    vapid_claims={"sub": PUSH_VAPID_CLAIMS_EMAIL},
                )
                connection.execute("""
                    UPDATE push_subscriptions
                    SET last_used_at = ?
                    WHERE id = ?
                """, (datetime.now().isoformat(timespec="seconds"), subscription["id"]))
                sent += 1
            except WebPushException as error:
                if getattr(error, "response", None) is not None and error.response.status_code in {404, 410}:
                    connection.execute(
                        "DELETE FROM push_subscriptions WHERE id = ?",
                        (subscription["id"],)
                    )
        connection.commit()
        return sent
    finally:
        connection.close()


def create_notification(user_id, title, body, link=None, kind="general", push=True):
    connection = get_db_connection()
    try:
        connection.execute("""
            INSERT INTO notifications (
                user_id, title, body, link, kind, created_at
            )
            VALUES (?, ?, ?, ?, ?, ?)
        """, (
            user_id, title, body, link, kind,
            datetime.now().isoformat(timespec="seconds")
        ))
        connection.commit()
    finally:
        connection.close()
    if push:
        send_push_to_user(user_id, title, body, link or "/notifications")

def get_unread_notification_count(user_id):
    connection = get_db_connection()
    try:
        return connection.execute("""
            SELECT COUNT(*) FROM notifications
            WHERE user_id = ? AND is_read = 0
        """, (user_id,)).fetchone()[0]
    finally:
        connection.close()


def login_required(view):
    @wraps(view)
    def wrapped_view(*args, **kwargs):
        if get_current_user() is None:
            if request.path.startswith("/api/"):
                return jsonify({
                    "error": "Please log in to use this feature."
                }), 401

            return redirect(url_for("login", next=request.path))

        return view(*args, **kwargs)

    return wrapped_view


def teacher_required(view):
    @wraps(view)
    @login_required
    def wrapped_view(*args, **kwargs):
        current_user = get_current_user()
        if not current_user or current_user["role"] != "teacher":
            return (
                render_template(
                    "access_denied.html",
                    message="Teacher mode is restricted to approved classroom accounts."
                ),
                403,
            )
        return view(*args, **kwargs)
    return wrapped_view


@app.context_processor
def inject_current_user():
    user = get_current_user()
    user_settings = get_user_settings(user["id"]) if user else None
    if user:
        sync_learning_notifications(user["id"])
    return {
        "current_user": user,
        "csrf_field": csrf_field,
        "profile_avatar_options": PROFILE_AVATAR_OPTIONS,
        "user_settings": user_settings,
        "theme": user_settings["theme"] if user_settings else "dark",
        "left_handed_mode": bool(
            user_settings["left_handed_mode"]
        ) if user_settings else False,
        "unread_notification_count": (
            get_unread_notification_count(user["id"]) if user else 0
        ),
        "vapid_public_key": PUSH_VAPID_PUBLIC_KEY,
    }


@app.before_request
def require_login_for_app_pages():
    public_endpoints = {"login", "register", "static"}

    if request.endpoint in public_endpoints:
        return None

    if get_current_user() is not None:
        return None

    if request.path.startswith("/api/"):
        return jsonify({
            "error": "Please log in to use this feature."
        }), 401

    return redirect(url_for("login", next=request.path))


def save_lecture(subject, title, content, user_id=None):
    user_id = user_id or get_current_user()["id"]
    connection = get_db_connection()

    try:
        connection.execute("""
            INSERT INTO lectures (user_id, subject, date, title, content)
            VALUES (?, ?, ?, ?, ?)
        """, (
            user_id,
            subject,
            date.today().strftime("%d %B %Y"),
            title,
            content
        ))

        connection.commit()

    finally:
        connection.close()


def save_quiz_attempt(
    quiz_type,
    lecture_id,
    lecture_title,
    score,
    max_score
):
    """Save a completed quiz attempt to quiz_history."""

    percentage = (
        round((score / max_score) * 100, 2)
        if max_score > 0
        else 0
    )

    attempted_at = datetime.now().strftime(
        "%Y-%m-%d %H:%M:%S"
    )

    connection = get_db_connection()

    try:
        connection.execute("""
            INSERT INTO quiz_history (
                user_id,
                quiz_type,
                lecture_id,
                lecture_title,
                score,
                max_score,
                percentage,
                attempted_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            get_current_user()["id"],
            quiz_type,
            lecture_id,
            lecture_title,
            score,
            max_score,
            percentage,
            attempted_at
        ))

        connection.commit()

    finally:
        connection.close()


def get_lecture_for_quiz(lecture_id):
    """Return a lecture row or None."""

    connection = get_db_connection()

    try:
        return connection.execute("""
            SELECT id, subject, date, title, content
            FROM lectures
            WHERE id = ? AND (user_id = ? OR user_id IS NULL)
        """, (lecture_id, get_current_user()["id"])).fetchone()

    finally:
        connection.close()


def parse_ai_json(result_text):
    """Parse JSON returned by the AI, including fenced JSON."""

    result_text = (result_text or "").strip()

    if result_text.startswith("```"):
        result_text = re.sub(
            r"^```(?:json)?\s*",
            "",
            result_text,
            count=1,
            flags=re.IGNORECASE
        )

        result_text = re.sub(
            r"\s*```\s*$",
            "",
            result_text,
            count=1
        )

    return json.loads(result_text)


# =========================================================
# AUTHENTICATION
# =========================================================

@app.route("/login", methods=["GET", "POST"])
def login():
    if get_current_user() is not None:
        return redirect(url_for("home"))

    if request.method == "POST":
        if auth_rate_limited("login"):
            return render_template(
                "auth.html",
                mode="login",
                error="Too many attempts. Please wait a few minutes and try again."
            ), 429
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")

        connection = get_db_connection()

        try:
            user = connection.execute(
                """SELECT id, name, username, email, password_hash
                   FROM users
                   WHERE lower(email) = ? OR lower(username) = ?""",
                (email, email)
            ).fetchone()
        finally:
            connection.close()

        if user is None or not check_password_hash(
            user["password_hash"],
            password
        ):
            return render_template(
                "auth.html",
                mode="login",
                error="Email or password is incorrect."
            )

        session.clear()
        session["user_id"] = user["id"]

        next_url = request.args.get("next", "")
        if next_url.startswith("/") and not next_url.startswith("//"):
            return redirect(next_url)

        return redirect(url_for("home"))

    return render_template("auth.html", mode="login")


@app.route("/register", methods=["GET", "POST"])
def register():
    if get_current_user() is not None:
        return redirect(url_for("home"))

    if request.method == "POST":
        if auth_rate_limited("register"):
            return render_template(
                "auth.html",
                mode="register",
                error="Too many attempts. Please wait a few minutes and try again."
            ), 429
        name = request.form.get("name", "").strip()
        username = request.form.get("username", "").strip().lower()
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        requested_role = request.form.get("role", "student").strip().lower()
        contacts_enabled = 1 if request.form.get("contacts_enabled") == "on" else 0

        if len(name) < 2:
            return render_template(
                "auth.html",
                mode="register",
                error="Please enter your name."
            )

        if not re.fullmatch(r"[a-z0-9_]{3,24}", username):
            return render_template(
                "auth.html",
                mode="register",
                error="Username must be 3–24 characters using lowercase letters, numbers, or underscores."
            )

        if "@" not in email or "." not in email.split("@")[-1]:
            return render_template(
                "auth.html",
                mode="register",
                error="Please enter a valid email address."
            )

        if len(password) < 8:
            return render_template(
                "auth.html",
                mode="register",
                error="Your password must contain at least 8 characters."
            )

        connection = get_db_connection()

        try:
            connection.execute("""
                INSERT INTO users (
                    name, username, email, password_hash, created_at, role
                )
                VALUES (?, ?, ?, ?, ?, ?)
            """, (
                name,
                username,
                email,
                generate_password_hash(password),
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "teacher" if requested_role == "teacher" and email in TEACHER_EMAILS else "student"
            ))
            connection.commit()
            user_id = connection.execute(
                "SELECT id FROM users WHERE email = ?",
                (email,)
            ).fetchone()["id"]
            connection.execute("""
                INSERT INTO user_settings (user_id, contacts_enabled)
                VALUES (?, ?)
            """, (user_id, contacts_enabled))
            connection.commit()
        except sqlite3.IntegrityError:
            return render_template(
                "auth.html",
                mode="register",
                error="That email or username is already in use."
            )
        finally:
            connection.close()

        session.clear()
        session["user_id"] = user_id
        get_user_settings(user_id)
        return redirect(url_for("home"))

    return render_template("auth.html", mode="register")


@app.route("/logout", methods=["POST"])
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/settings", methods=["GET", "POST"])
def settings():
    user = get_current_user()
    user_settings = get_user_settings(user["id"])

    if request.method == "POST":
        name = request.form.get("name", "").strip()
        email = request.form.get("email", "").strip().lower()
        language = request.form.get("language", "bilingual")
        notifications_enabled = int(
            request.form.get("notifications_enabled", "0") == "1"
        )
        focus_reminders_enabled = int(
            request.form.get("focus_reminders_enabled", "0") == "1"
        )
        contacts_enabled = int(
            request.form.get("contacts_enabled", "0") == "1"
        )
        theme = request.form.get("theme", "dark")
        left_handed_mode = int(
            request.form.get("left_handed_mode", "0") == "1"
        )
        profile_picture = request.files.get("profile_picture")
        selected_avatar = request.form.get("profile_avatar", "").strip()
        profile_picture_filename = user["profile_picture"]
        old_profile_picture = user["profile_picture"]

        if len(name) < 2:
            return render_template(
                "settings.html",
                user=user,
                user_settings=user_settings,
                error="Please enter a name with at least 2 characters."
            )

        if "@" not in email or "." not in email.split("@")[-1]:
            return render_template(
                "settings.html",
                user=user,
                user_settings=user_settings,
                error="Please enter a valid email address."
            )

        if selected_avatar and selected_avatar not in PROFILE_AVATAR_VALUES:
            selected_avatar = ""

        if selected_avatar:
            profile_picture_filename = selected_avatar
        elif profile_picture and profile_picture.filename:
            original_filename = secure_filename(profile_picture.filename)
            extension = original_filename.rsplit(".", 1)[-1].lower() if "." in original_filename else ""

            if extension not in ALLOWED_PROFILE_PICTURE_EXTENSIONS:
                return render_template(
                    "settings.html",
                    user=user,
                    user_settings=user_settings,
                    error="Please upload a JPG, PNG, WEBP, or GIF image."
                )

            profile_picture_filename = (
                f"user-{user['id']}-{uuid4().hex}.{extension}"
            )

        if language not in {"english", "tamil", "bilingual"}:
            language = "bilingual"
        if theme not in AVAILABLE_THEMES:
            theme = "dark"

        connection = get_db_connection()

        try:
            existing_user = connection.execute("""
                SELECT id FROM users
                WHERE email = ? AND id != ?
            """, (email, user["id"])).fetchone()

            if existing_user is not None:
                return render_template(
                    "settings.html",
                    user=user,
                    user_settings=user_settings,
                    error="That email address is already used by another account."
                )

            connection.execute("""
                UPDATE users
                SET name = ?, email = ?, profile_picture = ?
                WHERE id = ?
            """, (name, email, profile_picture_filename, user["id"]))
            connection.execute("""
                INSERT INTO user_settings (
                    user_id, language, notifications_enabled,
                    focus_reminders_enabled, theme, left_handed_mode,
                    contacts_enabled
                )
                VALUES (?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(user_id) DO UPDATE SET
                    language = excluded.language,
                    notifications_enabled = excluded.notifications_enabled,
                    focus_reminders_enabled = excluded.focus_reminders_enabled,
                    theme = excluded.theme,
                    left_handed_mode = excluded.left_handed_mode,
                    contacts_enabled = excluded.contacts_enabled
            """, (
                user["id"],
                language,
                notifications_enabled,
                focus_reminders_enabled,
                theme,
                left_handed_mode,
                contacts_enabled
            ))
            connection.commit()
        finally:
            connection.close()

        if profile_picture and profile_picture.filename and not selected_avatar:
            profile_picture.save(
                os.path.join(PROFILE_PICTURE_FOLDER, profile_picture_filename)
            )

        if (
            old_profile_picture
            and old_profile_picture != profile_picture_filename
            and not old_profile_picture.startswith("preset:")
        ):
            old_picture_path = os.path.join(
                PROFILE_PICTURE_FOLDER,
                old_profile_picture
            )
            if os.path.isfile(old_picture_path):
                os.remove(old_picture_path)

        return redirect(url_for("settings", saved="1"))

    return render_template(
        "settings.html",
        user=user,
        user_settings=user_settings,
        saved=request.args.get("saved") == "1"
    )


@app.route("/profile-pictures/<path:filename>")
@login_required
def profile_picture(filename):
    if filename.startswith("preset:"):
        abort(404)

    connection = get_db_connection()
    try:
        picture_exists = connection.execute(
            "SELECT 1 FROM users WHERE profile_picture = ?",
            (filename,)
        ).fetchone()
    finally:
        connection.close()

    if picture_exists is None:
        abort(404)

    return send_from_directory(PROFILE_PICTURE_FOLDER, filename)


def sync_learning_notifications(user_id):
    today = date.today().isoformat()
    connection = get_db_connection()
    try:
        quiz_today = connection.execute("""
            SELECT 1 FROM quiz_history
            WHERE user_id = ? AND attempted_at LIKE ? LIMIT 1
        """, (user_id, today + "%")).fetchone()
        existing_task = connection.execute("""
            SELECT 1 FROM notifications
            WHERE user_id = ? AND kind = 'daily-task' AND created_at LIKE ?
        """, (user_id, today + "%")).fetchone()
        if quiz_today is None and existing_task is None:
            connection.execute("""
                INSERT INTO notifications (
                    user_id, title, body, link, kind, created_at
                )
                VALUES (?, ?, ?, ?, 'daily-task', ?)
            """, (
                user_id,
                "Today’s learning task is waiting",
                "Complete one practice quiz or review a lecture before the day ends.",
                "/#test-understanding",
                datetime.now().isoformat(timespec="seconds")
            ))
            connection.commit()
    finally:
        connection.close()


@app.route("/notifications")
def notifications():
    user_id = get_current_user()["id"]
    sync_learning_notifications(user_id)
    connection = get_db_connection()
    try:
        items = connection.execute("""
            SELECT id, title, body, link, kind, is_read, created_at
            FROM notifications
            WHERE user_id = ?
            ORDER BY id DESC
            LIMIT 50
        """, (user_id,)).fetchall()
    finally:
        connection.close()
    return render_template("notifications.html", notifications=items)


@app.route("/notifications/read", methods=["POST"])
def mark_notifications_read():
    user_id = get_current_user()["id"]
    connection = get_db_connection()
    try:
        connection.execute(
            "UPDATE notifications SET is_read = 1 WHERE user_id = ?",
            (user_id,)
        )
        connection.commit()
    finally:
        connection.close()
    return redirect(url_for("notifications"))


@app.route("/api/activity", methods=["POST"])
def activity_heartbeat():
    payload = request.get_json(silent=True) or {}
    record_user_activity(
        get_current_user()["id"],
        payload.get("active_seconds", 0)
    )
    return jsonify({"ok": True})


@app.route("/api/push-subscription", methods=["POST"])
def save_push_subscription():
    subscription = request.get_json(silent=True) or {}
    endpoint = subscription.get("endpoint")
    if not endpoint or not subscription.get("keys"):
        return jsonify({"error": "Invalid push subscription."}), 400

    user_id = get_current_user()["id"]
    connection = get_db_connection()
    try:
        connection.execute("""
            INSERT INTO push_subscriptions (
                user_id, endpoint, subscription_json,
                created_at, last_used_at
            )
            VALUES (?, ?, ?, ?, ?)
            ON CONFLICT(endpoint) DO UPDATE SET
                user_id = excluded.user_id,
                subscription_json = excluded.subscription_json,
                last_used_at = excluded.last_used_at
        """, (
            user_id,
            endpoint,
            json.dumps(subscription),
            datetime.now().isoformat(timespec="seconds"),
            datetime.now().isoformat(timespec="seconds"),
        ))
        connection.commit()
    finally:
        connection.close()
    return jsonify({"ok": True})


def send_due_reminders():
    """Create and push one daily reminder for students inactive for 24 hours."""
    cutoff = datetime.now() - timedelta(hours=24)
    today = date.today().isoformat()
    connection = get_db_connection()
    reminded_users = []
    try:
        users = connection.execute("""
            SELECT u.id, u.name,
                   COALESCE(a.last_seen_at, u.created_at) AS last_seen_at
            FROM users u
            JOIN user_settings s ON s.user_id = u.id
            LEFT JOIN user_activity a ON a.user_id = u.id
            WHERE COALESCE(a.last_seen_at, u.created_at) < ?
              AND s.notifications_enabled = 1
        """, (cutoff.isoformat(timespec="seconds"),)).fetchall()
        sent = 0
        for user in users:
            exists = connection.execute("""
                SELECT 1 FROM notifications
                WHERE user_id = ? AND kind = 'inactive-reminder'
                  AND created_at LIKE ?
            """, (user["id"], today + "%")).fetchone()
            if exists:
                continue
            connection.execute("""
                INSERT INTO notifications (
                    user_id, title, body, link, kind, created_at
                )
                VALUES (?, ?, ?, ?, 'inactive-reminder', ?)
            """, (
                user["id"],
                "Your study space is ready",
                f"Hi {user['name']}, you have not checked in today. Complete one small learning task to keep moving.",
                "/#test-understanding",
                datetime.now().isoformat(timespec="seconds")
            ))
            reminded_users.append(user)
            sent += 1
        connection.commit()
    finally:
        connection.close()

    for user in reminded_users:
        send_push_to_user(
            user["id"],
            "Your study space is ready",
            "Complete one small learning task to keep moving.",
            "/#test-understanding"
        )
    return sent


@app.route("/notifications/timer", methods=["POST"])
def timer_notification():
    user_id = get_current_user()["id"]
    create_notification(
        user_id,
        "Focus timer complete",
        "Great work. Your 25-minute study block is finished. Take a short break or continue.",
        "/#study-room",
        "timer"
    )
    return jsonify({"ok": True})


def get_friendship_status(connection, user_id, other_user_id):
    friendship = connection.execute("""
        SELECT requester_id, addressee_id, status
        FROM friendships
        WHERE (requester_id = ? AND addressee_id = ?)
           OR (requester_id = ? AND addressee_id = ?)
        ORDER BY id DESC
        LIMIT 1
    """, (user_id, other_user_id, other_user_id, user_id)).fetchone()

    if friendship is None:
        return None

    if friendship["status"] == "accepted":
        return "friends"
    if (
        friendship["status"] == "pending"
        and friendship["addressee_id"] == user_id
    ):
        return "incoming"
    if (
        friendship["status"] == "pending"
        and friendship["requester_id"] == user_id
    ):
        return "outgoing"
    return friendship["status"]


def render_community(user_id, error=None, notice=None, query=""):
    connection = get_db_connection()

    try:
        students = connection.execute("""
            SELECT id, name, username, email, profile_picture
            FROM users
            WHERE id != ?
              AND (
                  ? = '' OR username LIKE ? OR name LIKE ?
              )
            ORDER BY name COLLATE NOCASE
        """, (user_id, query, f"%{query}%", f"%{query}%")).fetchall()
        students = [
            dict(student, friendship_status=get_friendship_status(
                connection, user_id, student["id"]
            ))
            for student in students
        ]

        friends = connection.execute("""
            SELECT u.id, u.name, u.email, u.profile_picture
            FROM users u
            JOIN friendships f
              ON f.status = 'accepted'
             AND (
                 (f.requester_id = ? AND f.addressee_id = u.id)
                 OR (f.addressee_id = ? AND f.requester_id = u.id)
             )
            ORDER BY u.name COLLATE NOCASE
        """, (user_id, user_id)).fetchall()

        incoming = connection.execute("""
            SELECT f.id, u.id AS user_id, u.name, u.profile_picture
            FROM friendships f
            JOIN users u ON u.id = f.requester_id
            WHERE f.addressee_id = ? AND f.status = 'pending'
            ORDER BY f.created_at DESC
        """, (user_id,)).fetchall()

        groups = connection.execute("""
            SELECT g.id, g.name, g.description, g.owner_id,
                   u.name AS owner_name,
                   COUNT(gm_all.user_id) AS member_count,
                   EXISTS(
                       SELECT 1
                       FROM study_group_members gm_me
                       WHERE gm_me.group_id = g.id AND gm_me.user_id = ?
                   ) AS is_member
            FROM study_groups g
            JOIN users u ON u.id = g.owner_id
            LEFT JOIN study_group_members gm_all ON gm_all.group_id = g.id
            GROUP BY g.id
            ORDER BY g.created_at DESC
        """, (user_id,)).fetchall()
    finally:
        connection.close()

    return render_template(
        "community.html",
        students=students,
        friends=friends,
        incoming_requests=incoming,
        groups=groups,
        error=error,
        notice=notice,
        query=query,
    )


@app.route("/community")
def community():
    query = request.args.get("q", "").strip()[:40]
    return render_community(get_current_user()["id"], query=query)


@app.route("/api/community/contact-discovery", methods=["POST"])
def contact_discovery():
    user_id = get_current_user()["id"]
    data = request.get_json(silent=True) or {}
    contacts = data.get("contacts", [])
    if not isinstance(contacts, list):
        return jsonify({"error": "Invalid contacts list."}), 400
    emails = sorted({
        str(contact).strip().lower()
        for contact in contacts[:50]
        if isinstance(contact, str) and "@" in contact
    })
    if not emails:
        return jsonify({"matches": []})
    placeholders = ",".join("?" for _ in emails)
    connection = get_db_connection()
    try:
        matches = connection.execute(f"""
            SELECT u.id, u.name, u.username, u.email
            FROM users u
            JOIN user_settings s ON s.user_id = u.id
            WHERE u.id != ? AND s.contacts_enabled = 1
              AND lower(email) IN ({placeholders})
            ORDER BY name COLLATE NOCASE
        """, (user_id, *emails)).fetchall()
    finally:
        connection.close()
    return jsonify({"matches": [dict(row) for row in matches]})


@app.route("/community/friends/<int:user_id>", methods=["POST"])
def send_friend_request(user_id):
    current_user = get_current_user()

    if user_id == current_user["id"]:
        return render_community(current_user["id"], error="You cannot add yourself.")

    connection = get_db_connection()
    try:
        target = connection.execute(
            "SELECT id FROM users WHERE id = ?", (user_id,)
        ).fetchone()
        if target is None:
            return render_community(
                current_user["id"], error="That student account was not found."
            )

        status = get_friendship_status(connection, current_user["id"], user_id)
        if status == "friends":
            return render_community(
                current_user["id"], notice="You are already friends with that student."
            )
        if status == "incoming":
            return render_community(
                current_user["id"],
                notice="That student has already sent you a request."
            )
        if status == "outgoing":
            return render_community(
                current_user["id"],
                notice="Friend request already sent."
            )

        if status == "rejected":
            connection.execute("""
                UPDATE friendships
                SET requester_id = ?, addressee_id = ?,
                    status = 'pending', created_at = ?
                WHERE (
                    requester_id = ? AND addressee_id = ?
                ) OR (
                    requester_id = ? AND addressee_id = ?
                )
            """, (
                current_user["id"], user_id,
                datetime.now().isoformat(timespec="seconds"),
                current_user["id"], user_id,
                user_id, current_user["id"]
            ))
        else:
            connection.execute("""
                INSERT INTO friendships (
                    requester_id, addressee_id, status, created_at
                )
                VALUES (?, ?, 'pending', ?)
            """, (
                current_user["id"],
                user_id,
                datetime.now().isoformat(timespec="seconds")
            ))
        connection.commit()
    finally:
        connection.close()

    create_notification(
        user_id,
        "New friend request",
        f"{current_user['name']} wants to connect with you.",
        "/community",
        "friend-request"
    )
    return redirect(url_for("community"))


@app.route("/community/requests/<int:friendship_id>/<action>", methods=["POST"])
def handle_friend_request(friendship_id, action):
    current_user = get_current_user()
    if action not in {"accept", "decline"}:
        return render_community(current_user["id"], error="That action is not available.")

    connection = get_db_connection()
    try:
        friendship = connection.execute("""
            SELECT id FROM friendships
            WHERE id = ? AND addressee_id = ? AND status = 'pending'
        """, (friendship_id, current_user["id"])).fetchone()
        if friendship is None:
            return render_community(
                current_user["id"], error="That friend request is no longer available."
            )

        connection.execute(
            "UPDATE friendships SET status = ? WHERE id = ?",
            ("accepted" if action == "accept" else "rejected", friendship_id)
        )
        connection.commit()
    finally:
        connection.close()

    return redirect(url_for("community"))


@app.route("/community/groups", methods=["POST"])
def create_group():
    current_user = get_current_user()
    name = request.form.get("name", "").strip()
    description = request.form.get("description", "").strip()

    if len(name) < 3 or len(name) > 80:
        return render_community(
            current_user["id"],
            error="Group names must be between 3 and 80 characters."
        )
    if len(description) > 240:
        return render_community(
            current_user["id"],
            error="Group descriptions must be 240 characters or fewer."
        )

    connection = get_db_connection()
    try:
        cursor = connection.execute("""
            INSERT INTO study_groups (
                name, description, owner_id, created_at
            )
            VALUES (?, ?, ?, ?)
        """, (
            name,
            description,
            current_user["id"],
            datetime.now().isoformat(timespec="seconds")
        ))
        connection.execute("""
            INSERT INTO study_group_members (group_id, user_id, joined_at)
            VALUES (?, ?, ?)
        """, (
            cursor.lastrowid,
            current_user["id"],
            datetime.now().isoformat(timespec="seconds")
        ))
        connection.commit()
    finally:
        connection.close()

    return redirect(url_for("community"))


@app.route("/community/groups/<int:group_id>/<action>", methods=["POST"])
def update_group_membership(group_id, action):
    current_user = get_current_user()
    if action not in {"join", "leave"}:
        return render_community(
            current_user["id"], error="That group action is not available."
        )

    connection = get_db_connection()
    try:
        group = connection.execute(
            "SELECT id FROM study_groups WHERE id = ?", (group_id,)
        ).fetchone()
        if group is None:
            return render_community(
                current_user["id"], error="That study group was not found."
            )

        if action == "join":
            connection.execute("""
                INSERT OR IGNORE INTO study_group_members (
                    group_id, user_id, joined_at
                )
                VALUES (?, ?, ?)
            """, (
                group_id,
                current_user["id"],
                datetime.now().isoformat(timespec="seconds")
            ))
        else:
            connection.execute("""
                DELETE FROM study_group_members
                WHERE group_id = ? AND user_id = ?
            """, (group_id, current_user["id"]))
        connection.commit()
    finally:
        connection.close()

    return redirect(url_for("community"))


def are_friends(connection, user_id, other_user_id):
    return connection.execute("""
        SELECT 1 FROM friendships
        WHERE status = 'accepted'
          AND (
              (requester_id = ? AND addressee_id = ?)
              OR (requester_id = ? AND addressee_id = ?)
          )
        LIMIT 1
    """, (user_id, other_user_id, other_user_id, user_id)).fetchone() is not None


@app.route("/chat/<int:user_id>", methods=["GET", "POST"])
def chat(user_id):
    current_user = get_current_user()
    connection = get_db_connection()

    try:
        other_user = connection.execute("""
            SELECT id, name, email, profile_picture
            FROM users
            WHERE id = ?
        """, (user_id,)).fetchone()

        if other_user is None or not are_friends(
            connection, current_user["id"], user_id
        ):
            return redirect(url_for("community"))

        if request.method == "POST":
            body = request.form.get("body", "").strip()
            if not body:
                error = "Write a message before sending."
            elif len(body) > 2000:
                error = "Messages must be 2,000 characters or fewer."
            else:
                connection.execute("""
                    INSERT INTO messages (
                        sender_id, receiver_id, body, created_at
                    )
                    VALUES (?, ?, ?, ?)
                """, (
                    current_user["id"],
                    user_id,
                    body,
                    datetime.now().isoformat(timespec="seconds")
                ))
                connection.commit()
                return redirect(url_for("chat", user_id=user_id))
        else:
            error = None

        messages = connection.execute("""
            SELECT m.id, m.sender_id, m.body, m.created_at,
                   u.name AS sender_name
            FROM messages m
            JOIN users u ON u.id = m.sender_id
            WHERE (m.sender_id = ? AND m.receiver_id = ?)
               OR (m.sender_id = ? AND m.receiver_id = ?)
            ORDER BY m.id
        """, (current_user["id"], user_id, user_id, current_user["id"])).fetchall()
    finally:
        connection.close()

    return render_template(
        "chat.html",
        other_user=other_user,
        messages=messages,
        error=error,
    )


@app.route("/group-chat/<int:group_id>", methods=["GET", "POST"])
def group_chat(group_id):
    current_user = get_current_user()
    connection = get_db_connection()

    try:
        group = connection.execute("""
            SELECT g.id, g.name, g.description,
                   COUNT(gm_all.user_id) AS member_count
            FROM study_groups g
            LEFT JOIN study_group_members gm_all ON gm_all.group_id = g.id
            JOIN study_group_members gm_me
              ON gm_me.group_id = g.id AND gm_me.user_id = ?
            WHERE g.id = ?
            GROUP BY g.id
        """, (current_user["id"], group_id)).fetchone()

        if group is None:
            return redirect(url_for("community"))

        if request.method == "POST":
            body = request.form.get("body", "").strip()
            if not body:
                error = "Write a message before sending."
            elif len(body) > 2000:
                error = "Messages must be 2,000 characters or fewer."
            else:
                connection.execute("""
                    INSERT INTO study_group_messages (
                        group_id, sender_id, body, created_at
                    )
                    VALUES (?, ?, ?, ?)
                """, (
                    group_id,
                    current_user["id"],
                    body,
                    datetime.now().isoformat(timespec="seconds")
                ))
                connection.commit()
                return redirect(url_for("group_chat", group_id=group_id))
        else:
            error = None

        messages = connection.execute("""
            SELECT gm.id, gm.sender_id, gm.body, gm.created_at,
                   u.name AS sender_name
            FROM study_group_messages gm
            JOIN users u ON u.id = gm.sender_id
            WHERE gm.group_id = ?
            ORDER BY gm.id
        """, (group_id,)).fetchall()
    finally:
        connection.close()

    return render_template(
        "group_chat.html",
        group=group,
        messages=messages,
        error=error,
    )


# =========================================================
# MAIN DASHBOARD
# =========================================================

@app.route("/planner", methods=["GET", "POST"])
def planner():
    user_id = get_current_user()["id"]
    connection = get_db_connection()
    error = None
    if request.method == "POST":
        subject = request.form.get("subject", "").strip()
        exam_date = request.form.get("exam_date", "").strip()
        units = request.form.get("units", "").strip()
        try:
            daily_minutes = max(10, min(240, int(request.form.get("daily_minutes", "25"))))
            confidence = max(1, min(5, int(request.form.get("confidence", "3"))))
            datetime.strptime(exam_date, "%Y-%m-%d")
        except (ValueError, TypeError):
            error = "Enter a valid exam date and study time."
        if not subject or not exam_date:
            error = "Add a subject and exam date."
        if not error:
            cursor = connection.execute("""
                INSERT INTO exam_plans (
                    user_id, subject, exam_date, units,
                    daily_minutes, confidence, created_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?)
            """, (user_id, subject, exam_date, units, daily_minutes, confidence,
                  datetime.now().isoformat(timespec="seconds")))
            plan_id = cursor.lastrowid
            unit_list = [
                item.strip() for item in units.replace("\n", ",").split(",")
                if item.strip()
            ]
            if not unit_list:
                unit_list = [subject]
            exam_day = datetime.strptime(exam_date, "%Y-%m-%d").date()
            available_days = max(1, (exam_day - date.today()).days)
            schedule_days = min(available_days, 60)
            created_at = datetime.now().isoformat(timespec="seconds")
            for day_offset in range(schedule_days):
                due_date = date.today() + timedelta(days=day_offset)
                if day_offset == schedule_days - 1:
                    title = f"Final review: {subject}"
                else:
                    unit = unit_list[day_offset % len(unit_list)]
                    title = f"Review {unit} ({daily_minutes} min)"
                connection.execute("""
                    INSERT INTO study_tasks (user_id, plan_id, title, due_date, created_at)
                    VALUES (?, ?, ?, ?, ?)
                """, (user_id, plan_id, title, due_date.isoformat(), created_at))
            connection.commit()
            create_notification(user_id, "Exam plan created",
                                f"Your {subject} schedule has {schedule_days} daily tasks.",
                                "/planner", "planner")
            return redirect(url_for("planner"))
    plans = connection.execute("""
        SELECT id, subject, exam_date, units, daily_minutes, confidence
        FROM exam_plans WHERE user_id = ? ORDER BY exam_date
    """, (user_id,)).fetchall()
    tasks = connection.execute("""
        SELECT id, title, due_date, completed FROM study_tasks
        WHERE user_id = ? ORDER BY completed, due_date, id DESC
    """, (user_id,)).fetchall()
    announcements = connection.execute("""
        SELECT ca.title, ca.body, ca.created_at, c.name AS classroom_name
        FROM classroom_announcements ca
        LEFT JOIN classrooms c ON c.id = ca.classroom_id
        WHERE ca.classroom_id IS NULL
           OR ca.classroom_id IN (
               SELECT classroom_id FROM classroom_members WHERE user_id = ?
           )
        ORDER BY ca.id DESC LIMIT 10
    """, (user_id,)).fetchall()
    classrooms = connection.execute("""
        SELECT c.id, c.name, c.join_code, c.teacher_id, u.name AS teacher_name
        FROM classrooms c
        JOIN users u ON u.id = c.teacher_id
        JOIN classroom_members cm ON cm.classroom_id = c.id
        WHERE cm.user_id = ?
        ORDER BY c.name
    """, (user_id,)).fetchall()
    connection.close()
    return render_template("planner.html", plans=plans, tasks=tasks,
                           announcements=announcements, classrooms=classrooms,
                           error=error)


@app.route("/classrooms/join", methods=["POST"])
@login_required
def join_classroom():
    current_user = get_current_user()
    if current_user["role"] != "student":
        flash("Only student accounts can join a classroom.")
        return redirect(url_for("planner"))

    join_code = request.form.get("join_code", "").strip().upper()
    if not re.fullmatch(r"[A-Z0-9]{6}", join_code):
        flash("Enter a valid six-character class code.")
        return redirect(url_for("planner"))

    connection = get_db_connection()
    try:
        classroom = connection.execute(
            "SELECT id, name FROM classrooms WHERE join_code = ?",
            (join_code,)
        ).fetchone()
        if classroom is None:
            flash("That classroom code was not found.")
        else:
            connection.execute("""
                INSERT OR IGNORE INTO classroom_members
                    (classroom_id, user_id, joined_at)
                VALUES (?, ?, ?)
            """, (
                classroom["id"],
                current_user["id"],
                datetime.now().isoformat(timespec="seconds")
            ))
            connection.commit()
            flash(f"You joined {classroom['name']}.")
    finally:
        connection.close()
    return redirect(url_for("planner"))


@app.route("/planner/tasks/<int:task_id>/toggle", methods=["POST"])
def toggle_study_task(task_id):
    user_id = get_current_user()["id"]
    connection = get_db_connection()
    connection.execute("""
        UPDATE study_tasks SET completed = CASE completed WHEN 1 THEN 0 ELSE 1 END
        WHERE id = ? AND user_id = ?
    """, (task_id, user_id))
    connection.commit()
    connection.close()
    return redirect(url_for("planner"))


@app.route("/search")
def smart_search():
    query = request.args.get("q", "").strip()
    connection = get_db_connection()
    if query:
        like = f"%{query}%"
        results = connection.execute("""
            SELECT id, subject, title, date, content
            FROM lectures
            WHERE (user_id = ? OR user_id IS NULL)
              AND (title LIKE ? OR subject LIKE ? OR content LIKE ?)
            ORDER BY id DESC
        """, (get_current_user()["id"], like, like, like)).fetchall()
    else:
        results = []
    connection.close()
    return render_template("search.html", query=query, results=results)


@app.route("/exam-mode")
def exam_mode():
    user_id = get_current_user()["id"]
    connection = get_db_connection()
    plan = connection.execute("""
        SELECT subject, exam_date, units, daily_minutes
        FROM exam_plans WHERE user_id = ? ORDER BY exam_date LIMIT 1
    """, (user_id,)).fetchone()
    tasks = connection.execute("""
        SELECT title, due_date, completed FROM study_tasks
        WHERE user_id = ? ORDER BY completed, due_date LIMIT 10
    """, (user_id,)).fetchall()
    connection.close()
    return render_template("exam_mode.html", plan=plan, tasks=tasks)


@app.route("/analytics")
def analytics():
    user_id = get_current_user()["id"]
    connection = get_db_connection()
    try:
        quiz_summary = connection.execute("""
            SELECT COUNT(*) AS attempts,
                   COALESCE(AVG(percentage), 0) AS average_score,
                   COALESCE(MAX(percentage), 0) AS best_score
            FROM quiz_history
            WHERE user_id = ?
        """, (user_id,)).fetchone()
        subject_scores = connection.execute("""
            SELECT lecture_title, COUNT(*) AS attempts,
                   ROUND(AVG(percentage)) AS score
            FROM quiz_history
            WHERE user_id = ?
            GROUP BY lecture_title
            ORDER BY score DESC
        """, (user_id,)).fetchall()
        activity = connection.execute("""
            SELECT total_active_seconds, last_seen_at
            FROM user_activity WHERE user_id = ?
        """, (user_id,)).fetchone()
        completed_tasks = connection.execute("""
            SELECT COUNT(*) FROM study_tasks
            WHERE user_id = ? AND completed = 1
        """, (user_id,)).fetchone()[0]
    finally:
        connection.close()
    return render_template(
        "analytics.html",
        quiz_summary=quiz_summary,
        subject_scores=subject_scores,
        activity=activity,
        completed_tasks=completed_tasks,
    )


@app.route("/teacher", methods=["GET", "POST"])
@teacher_required
def teacher_dashboard():
    user_id = get_current_user()["id"]
    error = None
    if request.method == "POST":
        action = request.form.get("action", "announcement")
        connection = get_db_connection()
        try:
            if action == "create_classroom":
                name = request.form.get("classroom_name", "").strip()
                if not name or len(name) > 100:
                    error = "Add a classroom name of 1–100 characters."
                else:
                    join_code = None
                    for _ in range(10):
                        candidate = uuid4().hex[:6].upper()
                        if not connection.execute(
                            "SELECT 1 FROM classrooms WHERE join_code = ?",
                            (candidate,)
                        ).fetchone():
                            join_code = candidate
                            break
                    if join_code is None:
                        error = "Could not create a unique class code. Try again."
                    else:
                        connection.execute("""
                            INSERT INTO classrooms
                                (teacher_id, name, join_code, created_at)
                            VALUES (?, ?, ?, ?)
                        """, (
                            user_id, name, join_code,
                            datetime.now().isoformat(timespec="seconds")
                        ))
                        connection.commit()
                        flash(f"Classroom created. Share code {join_code}.")
                        return redirect(url_for("teacher_dashboard"))
            else:
                title = request.form.get("title", "").strip()
                body = request.form.get("body", "").strip()
                classroom_id = request.form.get("classroom_id", "").strip()
                if not title or not body:
                    error = "Add both an announcement title and message."
                elif len(title) > 120 or len(body) > 1000:
                    error = "Keep the announcement within the displayed limits."
                elif classroom_id and not connection.execute(
                    "SELECT 1 FROM classrooms WHERE id = ? AND teacher_id = ?",
                    (classroom_id, user_id)
                ).fetchone():
                    error = "Choose one of your classrooms."
                else:
                    classroom_value = int(classroom_id) if classroom_id else None
                    connection.execute("""
                        INSERT INTO classroom_announcements
                            (author_id, classroom_id, title, body, created_at)
                        VALUES (?, ?, ?, ?, ?)
                    """, (
                        user_id, classroom_value, title, body,
                        datetime.now().isoformat(timespec="seconds")
                    ))
                    if classroom_value:
                        student_ids = [
                            row["user_id"] for row in connection.execute("""
                                SELECT user_id FROM classroom_members
                                WHERE classroom_id = ?
                            """, (classroom_value,)).fetchall()
                        ]
                    else:
                        student_ids = [
                            row["id"] for row in connection.execute(
                                "SELECT id FROM users WHERE role = 'student'"
                            ).fetchall()
                        ]
                    connection.commit()
                    for student_id in student_ids:
                        create_notification(
                            student_id,
                            "New classroom announcement",
                            title,
                            "/planner",
                            "classroom",
                        )
                    return redirect(url_for("teacher_dashboard"))
        finally:
            connection.close()

    connection = get_db_connection()
    try:
        announcements = connection.execute("""
            SELECT ca.title, ca.body, ca.created_at, c.name AS classroom_name
            FROM classroom_announcements ca
            LEFT JOIN classrooms c ON c.id = ca.classroom_id
            WHERE author_id = ? OR author_id IS NULL
            ORDER BY ca.id DESC LIMIT 20
        """, (user_id,)).fetchall()
        classrooms = connection.execute("""
            SELECT c.id, c.name, c.join_code, COUNT(cm.user_id) AS student_count
            FROM classrooms c
            LEFT JOIN classroom_members cm ON cm.classroom_id = c.id
            WHERE c.teacher_id = ?
            GROUP BY c.id
            ORDER BY c.created_at DESC
        """, (user_id,)).fetchall()
    finally:
        connection.close()
    return render_template(
        "teacher.html",
        announcements=announcements,
        classrooms=classrooms,
        error=error,
    )


@app.route("/")
def home():
    connection = get_db_connection()

    try:
        user_id = get_current_user()["id"]
        lecture_count = connection.execute(
            "SELECT COUNT(*) FROM lectures WHERE user_id = ? OR user_id IS NULL",
            (user_id,)
        ).fetchone()[0]
        quiz_stats = connection.execute("""
            SELECT COUNT(*) AS attempts,
                   COALESCE(AVG(percentage), 0) AS average_score
            FROM quiz_history
            WHERE user_id = ?
        """, (user_id,)).fetchone()
        latest_lecture = connection.execute("""
            SELECT subject, title, date
            FROM lectures
            WHERE user_id = ? OR user_id IS NULL
            ORDER BY id DESC
            LIMIT 1
        """, (user_id,)).fetchone()
        activity = connection.execute("""
            SELECT last_seen_at, total_active_seconds
            FROM user_activity
            WHERE user_id = ?
        """, (user_id,)).fetchone()
        recent_quiz = connection.execute("""
            SELECT percentage, lecture_title
            FROM quiz_history
            WHERE user_id = ?
            ORDER BY id DESC
            LIMIT 1
        """, (user_id,)).fetchone()
        weak_topics = connection.execute("""
            SELECT lecture_title, ROUND(AVG(percentage)) AS score
            FROM quiz_history
            WHERE user_id = ?
            GROUP BY lecture_title
            HAVING AVG(percentage) < 70
            ORDER BY score ASC
            LIMIT 3
        """, (user_id,)).fetchall()
        points = connection.execute("""
            SELECT COALESCE(SUM(CASE WHEN percentage >= 80 THEN 20 ELSE 10 END), 0)
            FROM quiz_history
            WHERE user_id = ?
        """, (user_id,)).fetchone()[0]
        upcoming_plan = connection.execute("""
            SELECT id, subject, exam_date, daily_minutes
            FROM exam_plans
            WHERE user_id = ?
            ORDER BY exam_date
            LIMIT 1
        """, (user_id,)).fetchone()
    finally:
        connection.close()

    active_minutes = round((activity["total_active_seconds"] if activity else 0) / 60)
    coach_status = (
        "You have not checked in yet. Start with one small study task."
        if not activity
        else f"You have spent about {active_minutes} active minutes in your workspace."
    )
    coach_next_step = (
        "Open your latest lecture notes for a 10-minute review."
        if latest_lecture
        else "Add your first lecture so I can build a study path with you."
    )
    if recent_quiz and recent_quiz["percentage"] < 60:
        coach_next_step = f"Review {recent_quiz['lecture_title']} and retry a short quiz."

    return render_template(
        "index.html",
        dashboard_stats={
            "lecture_count": lecture_count,
            "quiz_attempts": quiz_stats["attempts"],
            "average_score": round(quiz_stats["average_score"]),
            "latest_lecture": dict(latest_lecture) if latest_lecture else None,
        },
        personal_coach={
            "status": coach_status,
            "next_step": coach_next_step,
            "active_minutes": active_minutes,
            "last_seen": activity["last_seen_at"] if activity else None,
        },
        weak_topics=weak_topics,
        points=points,
        upcoming_plan=upcoming_plan
    )


# =========================================================
# UPLOAD AUDIO, TRANSCRIBE, AND GENERATE NOTES
# =========================================================

@app.route("/upload-audio", methods=["GET", "POST"])
def upload_audio():

    if request.method == "GET":
        return render_template("upload_audio.html")

    subject = request.form.get("subject", "").strip()
    title = request.form.get("title", "").strip()
    audio = request.files.get("audio")

    if not subject or not title or not audio or not audio.filename:
        return render_template(
            "upload_audio.html",
            error=(
                "Please enter the subject, lecture title, "
                "and select an audio file."
            )
        )

    if client is None:
        return render_template(
            "upload_audio.html",
            error="Groq API key is missing. Check your .env file."
        )

    original_name = secure_filename(audio.filename)

    extension = (
        original_name.rsplit(".", 1)[-1].lower()
        if "." in original_name
        else ""
    )

    if extension not in ALLOWED_EXTENSIONS:
        return render_template(
            "upload_audio.html",
            error="Unsupported file type. Choose a supported audio file."
        )

    temp_path = None

    try:
        with tempfile.NamedTemporaryFile(
            mode="wb",
            suffix="." + extension,
            dir=app.config["UPLOAD_FOLDER"],
            delete=False
        ) as temp_file:
            temp_path = temp_file.name
            audio.save(temp_file)

        if os.path.getsize(temp_path) == 0:
            return render_template(
                "upload_audio.html",
                error="The selected file is empty."
            )

        with open(temp_path, "rb") as audio_file:
            transcription = client.audio.transcriptions.create(
                file=(original_name, audio_file.read()),
                model="whisper-large-v3-turbo",
                response_format="text"
            )

        transcript = (
            transcription
            if isinstance(transcription, str)
            else getattr(transcription, "text", "")
        ).strip()

        if not transcript:
            return render_template(
                "upload_audio.html",
                error="No speech was detected. Try a clearer recording."
            )

        response = client.chat.completions.create(
            model="openai/gpt-oss-20b",
            messages=[
                {
                    "role": "system",
                    "content": (
                        "You are a college lecture note assistant. "
                        "Create accurate, well-organized study notes "
                        "from the transcript. Use Markdown headings, "
                        "bullet points, definitions, and a short summary. "
                        "Do not invent unsupported facts. Mark unclear "
                        "information rather than guessing."
                    )
                },
                {
                    "role": "user",
                    "content": (
                        f"Subject: {subject}\n"
                        f"Lecture title: {title}\n\n"
                        f"Transcript:\n{transcript}"
                    )
                }
            ],
            temperature=0.3
        )

        ai_notes = response.choices[0].message.content

        if not ai_notes:
            raise ValueError("The AI returned empty notes.")

        save_lecture(subject, title, ai_notes)

        return redirect(url_for("notes"))

    except Exception:
        app.logger.exception(
            "Audio transcription or note generation failed"
        )

        return render_template(
            "upload_audio.html",
            error=(
                "Audio processing failed. Check your internet connection, "
                "Groq API access, audio format, file size, and account limits. "
                "See the VS Code terminal for details."
            )
        )

    finally:
        if temp_path and os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except OSError:
                app.logger.warning(
                    "Could not remove temporary audio file."
                )


# =========================================================
# VIEW SAVED LECTURES
# =========================================================

@app.route("/notes")
def notes():
    connection = get_db_connection()

    try:
        all_notes = connection.execute(
            "SELECT * FROM lectures WHERE user_id = ? OR user_id IS NULL ORDER BY id DESC",
            (get_current_user()["id"],)
        ).fetchall()

    finally:
        connection.close()

    formatted_notes = []

    for note in all_notes:
        note_data = dict(note)

        note_data["formatted_content"] = markdown.markdown(
            note["content"],
            extensions=["tables", "fenced_code"]
        )

        formatted_notes.append(note_data)

    return render_template(
        "notes.html",
        all_notes=formatted_notes
    )


# =========================================================
# ADD A LECTURE MANUALLY
# =========================================================

@app.route("/add-lecture", methods=["GET", "POST"])
def add_lecture():

    if request.method == "POST":
        subject = request.form.get("subject", "").strip()
        title = request.form.get("title", "").strip()
        content = request.form.get("content", "").strip()

        if subject and title and content:
            save_lecture(subject, title, content)
            return redirect(url_for("notes"))

        return render_template(
            "add_lecture.html",
            error="Please fill in the subject, title, and notes."
        )

    return render_template("add_lecture.html")


# =========================================================
# GENERATE NOTES FROM A PASTED TRANSCRIPT
# =========================================================

@app.route("/generate-notes", methods=["GET", "POST"])
def generate_notes():

    if request.method == "POST":
        subject = request.form.get("subject", "").strip()
        title = request.form.get("title", "").strip()
        transcript = request.form.get("transcript", "").strip()
        source_file = request.files.get("source_file")
        language = request.form.get("language", "english").strip().lower()

        language_instruction = {
            "english": "Write the notes in clear English.",
            "tamil": "Write the notes in clear Tamil. Keep important English technical terms in parentheses.",
            "bilingual": "Write each important idea first in English and then explain it in Tamil. Keep technical terms in English where useful.",
        }.get(language, "Write the notes in clear English.")

        if source_file and source_file.filename:
            original_name = secure_filename(source_file.filename)
            extension = (
                original_name.rsplit(".", 1)[-1].lower()
                if "." in original_name
                else ""
            )
            if extension not in ALLOWED_DOCUMENT_EXTENSIONS:
                return render_template(
                    "generate_notes.html",
                    error="Unsupported document type. Use PDF, DOCX, PPTX, TXT, or MD."
                )
            try:
                transcript = extract_document_text(source_file, extension)
            except Exception:
                app.logger.exception("Document text extraction failed")
                return render_template(
                    "generate_notes.html",
                    error="We could not read that file. Check that it is a valid, text-based document."
                )

        if not subject or not title or not transcript:
            return render_template(
                "generate_notes.html",
                error="Please enter a subject, title, and either paste text or upload a document."
            )

        if client is None:
            return render_template(
                "generate_notes.html",
                error="Groq API key is missing. Check your .env file."
            )

        try:
            response = client.chat.completions.create(
                model="openai/gpt-oss-20b",
                messages=[
                    {
                        "role": "system",
                        "content": (
                            "You are a college lecture note assistant. "
                            "Create accurate, well-organized study notes "
                            "from the supplied transcript. Use headings, "
                            "bullet points, definitions, and a summary. "
                            "Do not invent unsupported facts. "
                            + language_instruction
                        )
                    },
                    {
                        "role": "user",
                        "content": (
                            f"Subject: {subject}\n"
                            f"Lecture title: {title}\n\n"
                            f"Transcript:\n{transcript}"
                        )
                    }
                ],
                temperature=0.3
            )

            ai_notes = response.choices[0].message.content

            if not ai_notes:
                raise ValueError("The AI returned empty notes.")

            save_lecture(subject, title, ai_notes)

            return redirect(url_for("notes"))

        except Exception:
            app.logger.exception("AI note generation failed")

            return render_template(
                "generate_notes.html",
                error=(
                    "AI note generation failed. Check your connection, "
                    "API key, model availability, and account limits."
                )
            )

    return render_template("generate_notes.html")


# =========================================================
# SMART LECTURE AI AGENT PAGE
# =========================================================

@app.route("/agent", methods=["GET", "POST"])
def agent():
    answer = None
    sources = []
    question = ""

    if request.method == "POST":
        question = request.form.get("question", "").strip()

        if not question:
            answer = "Please enter a question."

        elif client is None:
            answer = "Groq API key is missing. Check your .env file."

        else:
            connection = get_db_connection()

            try:
                lectures = connection.execute("""
                    SELECT id, subject, date, title, content
                    FROM lectures
                    WHERE user_id = ? OR user_id IS NULL
                    ORDER BY id DESC
                """, (get_current_user()["id"],)).fetchall()

            finally:
                connection.close()

            if not lectures:
                answer = (
                    "No saved lectures found. Add lecture notes first."
                )

            else:
                lecture_context_parts = []

                for lecture in lectures:
                    lecture_context_parts.append(
                        f"Lecture ID: {lecture['id']}\n"
                        f"Subject: {lecture['subject']}\n"
                        f"Date: {lecture['date']}\n"
                        f"Title: {lecture['title']}\n"
                        f"Notes:\n{lecture['content']}"
                    )

                lecture_context = "\n\n" + (
                    "\n" + ("-" * 40) + "\n"
                ).join(lecture_context_parts)

                try:
                    response = client.chat.completions.create(
                        model="openai/gpt-oss-20b",
                        messages=[
                            {
                                "role": "system",
                                "content": (
                                    "You are Smart Lecture AI, a college "
                                    "study assistant. Use these exact sections:\n\n"
                                    "## 1. FROM MY NOTES\n"
                                    "Only information supported by the notes.\n\n"
                                    "## 2. ADDITIONAL INFORMATION\n"
                                    "Clearly label general knowledge. Write "
                                    "'None' if unnecessary.\n\n"
                                    "## 3. SOURCES USED\n"
                                    "List the exact supporting lecture title, "
                                    "subject, and date. Never invent sources. "
                                    "Say clearly when the notes are insufficient."
                                )
                            },
                            {
                                "role": "user",
                                "content": (
                                    f"Saved lecture notes:\n{lecture_context}\n\n"
                                    f"Student question: {question}"
                                )
                            }
                        ],
                        temperature=0.3
                    )

                    answer = response.choices[0].message.content

                    if not answer:
                        answer = "The AI returned an empty answer."

                    sources = [
                        {
                            "subject": lecture["subject"],
                            "title": lecture["title"],
                            "date": lecture["date"]
                        }
                        for lecture in lectures
                    ]

                except Exception:
                    app.logger.exception("AI agent failed")
                    answer = (
                        "The AI could not answer right now. Check your "
                        "internet connection, Groq API access, and terminal."
                    )

    return render_template(
        "agent.html",
        answer=answer,
        sources=sources,
        question=question
    )


# =========================================================
# DASHBOARD AI AGENT API
# Supports notes, web search, both, and conversation history
# =========================================================

@app.route("/api/agent", methods=["POST"])
def dashboard_agent():
    data = request.get_json(silent=True) or {}

    question = data.get("question", "")

    if not isinstance(question, str):
        question = ""

    question = question.strip()

    source_mode = data.get("source", "notes")

    if not isinstance(source_mode, str):
        source_mode = "notes"

    source_mode = source_mode.strip().lower()

    history = data.get("history", [])
    valid_history = []

    if isinstance(history, list):
        for message in history[-10:]:
            if not isinstance(message, dict):
                continue

            role = message.get("role")
            content = message.get("content")

            if role not in {"user", "assistant"}:
                continue

            if not isinstance(content, str):
                continue

            content = content.strip()

            if content:
                valid_history.append({
                    "role": role,
                    "content": content[:4000]
                })

    if not question:
        return jsonify({
            "error": "Please enter a question."
        }), 400

    if source_mode not in {"notes", "web", "both"}:
        return jsonify({
            "error": "Invalid information source."
        }), 400

    if client is None:
        return jsonify({
            "error": "Groq API key is missing. Check your .env file."
        }), 500

    lecture_context = ""

    if source_mode in {"notes", "both"}:
        connection = get_db_connection()

        try:
            lectures = connection.execute("""
                SELECT subject, date, title, content
                FROM lectures
                WHERE user_id = ? OR user_id IS NULL
                ORDER BY id DESC
            """, (get_current_user()["id"],)).fetchall()

        finally:
            connection.close()

        if not lectures and source_mode == "notes":
            return jsonify({
                "error": (
                    "No saved lectures found. Add lecture notes first."
                )
            }), 404

        lecture_context = "\n\n".join(
            f"Subject: {lecture['subject']}\n"
            f"Date: {lecture['date']}\n"
            f"Title: {lecture['title']}\n"
            f"Notes:\n{lecture['content']}"
            for lecture in lectures
        )

    web_results = []

    if source_mode in {"web", "both"}:
        previous_questions = [
            message["content"]
            for message in valid_history
            if message["role"] == "user"
        ]

        search_query = question

        if previous_questions and len(question.split()) <= 8:
            previous_context = " ".join(previous_questions[-3:])
            search_query = f"{previous_context} {question}"[:500]

        try:
            with DDGS() as search_engine:
                results = search_engine.text(
                    search_query,
                    max_results=5
                )

                for result in results:
                    url = result.get("href", "")

                    if not url.startswith(("https://", "http://")):
                        continue

                    web_results.append({
                        "title": result.get("title", "") or url,
                        "url": url,
                        "snippet": result.get("body", "")
                    })

        except Exception:
            app.logger.exception("Internet search failed")

            return jsonify({
                "error": (
                    "Internet search failed. Check your connection. "
                    "The search service may be temporarily unavailable."
                )
            }), 502

        if not web_results:
            return jsonify({
                "error": (
                    "No useful Internet results were found. "
                    "Try a more specific question."
                )
            }), 502

    if source_mode == "notes":
        system_prompt = """
You are Smart Lecture AI, a college study assistant.

Answer using ONLY the supplied saved lecture notes.
Do not use outside knowledge to fill gaps.

## 1. FROM MY NOTES
Give the answer supported by the notes.

## 2. ADDITIONAL INFORMATION
Write "Not included because My Notes mode was selected."

## 3. SOURCES USED
List only supporting lecture titles, subjects, and dates.
Never invent source details.
"""

    elif source_mode == "web":
        system_prompt = """
You are Smart Lecture AI, a college study assistant.

Answer using the supplied Internet search results.
Do not claim to have read webpages beyond the supplied results.

## 1. FROM MY NOTES
Write "Not included because Internet mode was selected."

## 2. ADDITIONAL INFORMATION FROM THE INTERNET
Explain the answer using the supplied results.

## 3. SOURCES USED
List relevant webpage titles and exact URLs supplied.
Never invent URLs or citations.
Say if the evidence is insufficient.
"""

    else:
        system_prompt = """
You are Smart Lecture AI, a college study assistant.

Use both saved lecture notes and Internet search results.

## 1. FROM MY NOTES
Include only claims supported by saved notes.
Cite the relevant lecture title, subject, and date.

## 2. ADDITIONAL INFORMATION FROM THE INTERNET
Include information supported by the supplied search results.
Do not present Internet information as if it came from lecture notes.

## 3. SOURCES USED
List only sources that support the answer.
Never invent source details or citations.
Explain when information is insufficient.
"""

    messages = [
        {
            "role": "system",
            "content": system_prompt
        }
    ]

    if source_mode in {"notes", "both"}:
        messages.append({
            "role": "system",
            "content": (
                "Saved lecture notes:\n"
                + (lecture_context or "No saved lectures.")
            )
        })

    if source_mode in {"web", "both"}:
        web_context = "\n\n".join(
            f"Title: {item['title']}\n"
            f"URL: {item['url']}\n"
            f"Summary: {item['snippet']}"
            for item in web_results
        )

        messages.append({
            "role": "system",
            "content": "Internet search results:\n" + web_context
        })

    messages.extend(valid_history)

    messages.append({
        "role": "user",
        "content": question
    })

    try:
        response = client.chat.completions.create(
            model="openai/gpt-oss-20b",
            messages=messages,
            temperature=0.3
        )

        answer = response.choices[0].message.content

        if not answer:
            return jsonify({
                "error": "The AI returned an empty answer."
            }), 502

        frontend_sources = [
            {
                "title": item["title"],
                "href": item["url"],
                "body": item["snippet"]
            }
            for item in web_results
        ]

        return jsonify({
            "answer": answer,
            "source_mode": source_mode,
            "web_sources": frontend_sources
        })

    except Exception:
        app.logger.exception("Dashboard AI Agent failed")

        return jsonify({
            "error": (
                "The AI could not answer. Check your Groq API access, "
                "internet connection, and VS Code terminal."
            )
        }), 500


@app.route("/api/personal-coach", methods=["POST"])
def personal_coach():
    user_id = get_current_user()["id"]
    data = request.get_json(silent=True) or {}
    message = data.get("message", "")
    if not isinstance(message, str):
        message = ""
    message = message.strip()[:1000]

    connection = get_db_connection()
    try:
        activity = connection.execute("""
            SELECT total_active_seconds FROM user_activity
            WHERE user_id = ?
        """, (user_id,)).fetchone()
        quiz_stats = connection.execute("""
            SELECT COUNT(*) AS attempts,
                   COALESCE(AVG(percentage), 0) AS average_score
            FROM quiz_history
            WHERE user_id = ?
        """, (user_id,)).fetchone()
        pending_tasks = connection.execute("""
            SELECT COUNT(*) FROM notifications
            WHERE user_id = ? AND kind IN ('daily-task', 'inactive-reminder')
              AND is_read = 0
        """, (user_id,)).fetchone()[0]
        latest = connection.execute("""
            SELECT title FROM lectures
            WHERE user_id = ? OR user_id IS NULL
            ORDER BY id DESC LIMIT 1
        """, (user_id,)).fetchone()
    finally:
        connection.close()

    active_minutes = round((activity["total_active_seconds"] if activity else 0) / 60)
    average_score = round(quiz_stats["average_score"])
    if message:
        response = (
            f"I’m with you. You have {active_minutes} active study minutes and "
            f"an average quiz score of {average_score}%. Try one focused "
            "25-minute block, then tell me what felt difficult."
        )
    elif pending_tasks:
        response = "You have a pending learning task. Let’s complete one quiz or review one lecture now."
    elif average_score and average_score < 60:
        response = "Your recent quiz performance suggests a review block would help. Start with your weakest lecture topic."
    elif latest:
        response = f"Your next best step is a 10-minute review of “{latest['title']}”, followed by a short quiz."
    else:
        response = "Add a lecture and I’ll help turn it into a personal study path."

    return jsonify({
        "message": response,
        "active_minutes": active_minutes,
        "average_score": average_score,
        "pending_tasks": pending_tasks,
        "privacy": "I use your Smart Lecture AI activity and learning results. I do not read private device content."
    })


# =========================================================
# MCQ QUIZ PAGE
# =========================================================

@app.route("/quiz/mcqs")
def mcq_quiz():
    connection = get_db_connection()

    try:
        lectures = connection.execute("""
            SELECT id, subject, date, title
            FROM lectures
            WHERE user_id = ? OR user_id IS NULL
            ORDER BY id DESC
        """, (get_current_user()["id"],)).fetchall()

    finally:
        connection.close()

    return render_template(
        "quiz_mcqs.html",
        lectures=lectures
    )


# =========================================================
# GENERATE MCQ QUESTIONS
# =========================================================

@app.route("/api/quiz/generate", methods=["POST"])
def generate_mcq_quiz():
    if client is None:
        return jsonify({
            "error": "AI is not configured. Check your .env file."
        }), 500

    data = request.get_json(silent=True) or {}

    try:
        lecture_id = int(data.get("lecture_id", 0))
        question_count = int(data.get("question_count", 5))

    except (TypeError, ValueError):
        return jsonify({
            "error": "Invalid quiz settings."
        }), 400

    if lecture_id <= 0:
        return jsonify({
            "error": "Please select a valid lecture."
        }), 400

    if not 1 <= question_count <= 10:
        return jsonify({
            "error": "Choose between 1 and 10 questions."
        }), 400

    lecture = get_lecture_for_quiz(lecture_id)

    if lecture is None:
        return jsonify({
            "error": "Lecture not found."
        }), 404

    prompt = f"""
Create exactly {question_count} multiple-choice questions based ONLY
on the lecture content below.

Subject: {lecture['subject']}
Lecture title: {lecture['title']}

Lecture content:
{lecture['content']}

Return valid JSON in this exact structure:
{{
  "questions": [
    {{
      "question": "Question text",
      "options": ["Option A", "Option B", "Option C", "Option D"],
      "correct_index": 0,
      "explanation": "Short explanation based on the lecture"
    }}
  ]
}}

Rules:
- Each question must have exactly four options.
- correct_index must be 0, 1, 2, or 3.
- Only one option should be correct.
- Do not invent facts.
- Return JSON only.
"""

    try:
        response = client.chat.completions.create(
            model="openai/gpt-oss-20b",
            messages=[
                {
                    "role": "system",
                    "content": "You create accurate educational quizzes."
                },
                {
                    "role": "user",
                    "content": prompt
                }
            ],
            temperature=0.3
        )

        result = parse_ai_json(
            response.choices[0].message.content
        )

        questions = result.get("questions", [])

        if (
            not isinstance(questions, list)
            or len(questions) != question_count
        ):
            raise ValueError("Unexpected number of questions.")

        for question in questions:
            if not isinstance(question, dict):
                raise ValueError("Invalid question format.")

            options = question.get("options")
            correct_index = question.get("correct_index")

            if (
                not isinstance(question.get("question"), str)
                or not isinstance(options, list)
                or len(options) != 4
                or not all(isinstance(item, str) for item in options)
                or type(correct_index) is not int
                or correct_index not in range(4)
                or not isinstance(question.get("explanation"), str)
            ):
                raise ValueError("Invalid question data.")

        return jsonify({
            "questions": questions
        })

    except Exception:
        app.logger.exception("MCQ generation failed")

        return jsonify({
            "error": "Could not generate the quiz. Please try again."
        }), 500


# =========================================================
# SUBMIT MCQ QUIZ AND SAVE HISTORY
# =========================================================

@app.route("/api/quiz/submit", methods=["POST"])
def submit_mcq_quiz():
    data = request.get_json(silent=True) or {}

    questions = data.get("questions")
    answers = data.get("answers")
    lecture_id = data.get("lecture_id")

    if (
        not isinstance(questions, list)
        or not isinstance(answers, list)
    ):
        return jsonify({
            "error": "Invalid quiz submission."
        }), 400

    if (
        not 1 <= len(questions) <= 10
        or len(answers) != len(questions)
    ):
        return jsonify({
            "error": "Invalid number of answers."
        }), 400

    if type(lecture_id) is not int or lecture_id <= 0:
        return jsonify({
            "error": "Please select a valid lecture."
        }), 400

    lecture = get_lecture_for_quiz(lecture_id)

    if lecture is None:
        return jsonify({
            "error": "Lecture not found."
        }), 404

    score = 0
    results = []

    for question, answer in zip(questions, answers):
        if not isinstance(question, dict):
            return jsonify({
                "error": "Invalid question data."
            }), 400

        correct_index = question.get("correct_index")
        options = question.get("options")
        question_text = question.get("question", "")
        explanation = question.get("explanation", "")

        if (
            type(correct_index) is not int
            or correct_index not in range(4)
            or not isinstance(options, list)
            or len(options) != 4
            or not isinstance(question_text, str)
            or not isinstance(explanation, str)
        ):
            return jsonify({
                "error": "Invalid question data."
            }), 400

        if answer is not None and (
            type(answer) is not int or answer not in range(4)
        ):
            return jsonify({
                "error": "Invalid selected answer."
            }), 400

        is_correct = (
            type(answer) is int
            and answer == correct_index
        )

        if is_correct:
            score += 1

        results.append({
            "correct": is_correct,
            "selected_index": answer,
            "correct_index": correct_index,
            "question": question_text,
            "options": options,
            "explanation": explanation
        })

    total = len(questions)

    save_quiz_attempt(
        "MCQ",
        lecture_id,
        lecture["title"],
        score,
        total
    )

    return jsonify({
        "score": score,
        "total": total,
        "percentage": round((score / total) * 100, 2),
        "results": results
    })


# =========================================================
# FILL IN THE BLANKS QUIZ PAGE
# =========================================================

@app.route("/quiz/fill-blanks")
def fill_blanks_quiz():
    connection = get_db_connection()

    try:
        lectures = connection.execute("""
            SELECT id, subject, date, title
            FROM lectures
            WHERE user_id = ? OR user_id IS NULL
            ORDER BY id DESC
        """, (get_current_user()["id"],)).fetchall()

    finally:
        connection.close()

    return render_template(
        "quiz_fill_blanks.html",
        lectures=lectures
    )


# =========================================================
# GENERATE FILL IN THE BLANKS QUESTIONS
# =========================================================

@app.route("/api/quiz/fill-blanks/generate", methods=["POST"])
def generate_fill_blanks_quiz():
    if client is None:
        return jsonify({
            "error": "AI is not configured. Check your .env file."
        }), 500

    data = request.get_json(silent=True) or {}

    try:
        lecture_id = int(data.get("lecture_id", 0))
        question_count = int(data.get("question_count", 5))

    except (TypeError, ValueError):
        return jsonify({
            "error": "Invalid quiz settings."
        }), 400

    if lecture_id <= 0:
        return jsonify({
            "error": "Please select a valid lecture."
        }), 400

    if not 1 <= question_count <= 10:
        return jsonify({
            "error": "Choose between 1 and 10 questions."
        }), 400

    lecture = get_lecture_for_quiz(lecture_id)

    if lecture is None:
        return jsonify({
            "error": "Lecture not found."
        }), 404

    prompt = f"""
Create exactly {question_count} fill-in-the-blank questions based
ONLY on this lecture.

Subject: {lecture['subject']}
Title: {lecture['title']}

Lecture:
{lecture['content']}

Return JSON only:
{{
  "questions": [
    {{
      "sentence": "A resistor limits the flow of ____.",
      "answer": "electric current",
      "explanation": "A resistor limits the flow of electric current."
    }}
  ]
}}

Rules:
- Each sentence must contain exactly one ____ blank.
- Answers must be short words or phrases.
- Answers must be supported by the lecture.
- Avoid duplicate questions.
- Return exactly the requested number.
"""

    try:
        response = client.chat.completions.create(
            model="openai/gpt-oss-20b",
            messages=[
                {
                    "role": "system",
                    "content": (
                        "Create accurate fill-in-the-blank quizzes. "
                        "Return valid JSON only."
                    )
                },
                {
                    "role": "user",
                    "content": prompt
                }
            ],
            temperature=0.3
        )

        result = parse_ai_json(
            response.choices[0].message.content
        )

        questions = result.get("questions", [])

        if (
            not isinstance(questions, list)
            or len(questions) != question_count
        ):
            raise ValueError("Unexpected number of questions.")

        for question in questions:
            if not isinstance(question, dict):
                raise ValueError("Invalid question format.")

            sentence = question.get("sentence")
            answer = question.get("answer")
            explanation = question.get("explanation")

            if not all(
                isinstance(value, str) and value.strip()
                for value in (sentence, answer, explanation)
            ):
                raise ValueError("Missing question information.")

            if sentence.count("____") != 1:
                raise ValueError("Each sentence must contain one blank.")

        return jsonify({
            "questions": questions
        })

    except Exception:
        app.logger.exception(
            "Fill-in-the-blanks generation failed"
        )

        return jsonify({
            "error": "Could not generate questions. Please try again."
        }), 500


# =========================================================
# SUBMIT FILL IN THE BLANKS AND SAVE HISTORY
# =========================================================

@app.route("/api/quiz/fill-blanks/submit", methods=["POST"])
def submit_fill_blanks_quiz():
    data = request.get_json(silent=True) or {}

    questions = data.get("questions")
    answers = data.get("answers")
    lecture_id = data.get("lecture_id")

    if (
        not isinstance(questions, list)
        or not isinstance(answers, list)
        or len(questions) != len(answers)
        or not 1 <= len(questions) <= 10
    ):
        return jsonify({
            "error": "Invalid quiz submission."
        }), 400

    try:
        lecture_id = int(lecture_id)
    except (TypeError, ValueError):
        return jsonify({
            "error": "Please select a valid lecture."
        }), 400

    if lecture_id <= 0:
        return jsonify({
            "error": "Please select a valid lecture."
        }), 400

    lecture = get_lecture_for_quiz(lecture_id)

    if lecture is None:
        return jsonify({
            "error": "Lecture not found."
        }), 404

    score = 0
    results = []

    def normalize(text):
        return re.sub(r"\s+", " ", text.strip()).casefold()

    for question, answer in zip(questions, answers):
        if not isinstance(question, dict) or not isinstance(answer, str):
            return jsonify({
                "error": "Invalid answer data."
            }), 400

        correct_answer = question.get("answer")
        explanation = question.get("explanation", "")

        if (
            not isinstance(correct_answer, str)
            or not isinstance(explanation, str)
        ):
            return jsonify({
                "error": "Invalid question data."
            }), 400

        is_correct = normalize(answer) == normalize(correct_answer)

        if is_correct:
            score += 1

        results.append({
            "correct": is_correct,
            "student_answer": answer,
            "correct_answer": correct_answer,
            "explanation": explanation
        })

    total = len(questions)

    save_quiz_attempt(
        "Fill in the Blanks",
        lecture_id,
        lecture["title"],
        score,
        total
    )

    return jsonify({
        "score": score,
        "total": total,
        "percentage": round((score / total) * 100, 2),
        "results": results
    })


# =========================================================
# SHORT ANSWERS QUIZ PAGE
# =========================================================

@app.route("/quiz/short-answers")
def short_answers_quiz():
    connection = get_db_connection()

    try:
        lectures = connection.execute("""
            SELECT id, subject, date, title
            FROM lectures
            WHERE user_id = ? OR user_id IS NULL
            ORDER BY id DESC
        """, (get_current_user()["id"],)).fetchall()

    finally:
        connection.close()

    return render_template(
        "quiz_short_answers.html",
        lectures=lectures
    )


# =========================================================
# GENERATE SHORT ANSWER QUESTIONS
# =========================================================

@app.route("/api/quiz/short-answers/generate", methods=["POST"])
def generate_short_answers_quiz():
    if client is None:
        return jsonify({
            "error": "AI is not configured. Check your .env file."
        }), 500

    data = request.get_json(silent=True) or {}

    try:
        lecture_id = int(data.get("lecture_id", 0))
        question_count = int(data.get("question_count", 5))

    except (TypeError, ValueError):
        return jsonify({
            "error": "Invalid quiz settings."
        }), 400

    if lecture_id <= 0:
        return jsonify({
            "error": "Please select a valid lecture."
        }), 400

    if not 1 <= question_count <= 10:
        return jsonify({
            "error": "Choose between 1 and 10 questions."
        }), 400

    lecture = get_lecture_for_quiz(lecture_id)

    if lecture is None:
        return jsonify({
            "error": "Lecture not found."
        }), 404

    prompt = f"""
Create exactly {question_count} short-answer questions based ONLY
on this lecture.

Subject: {lecture['subject']}
Title: {lecture['title']}

Lecture:
{lecture['content']}

Return JSON only:
{{
  "questions": [
    {{
      "question": "What is the function of a resistor?",
      "expected_answer": "A resistor limits the flow of electric current."
    }}
  ]
}}

Rules:
- Questions must be answerable from the lecture.
- Expected answers should be concise and accurate.
- Avoid duplicates.
- Do not invent facts.
- Return exactly the requested number.
"""

    try:
        response = client.chat.completions.create(
            model="openai/gpt-oss-20b",
            messages=[
                {
                    "role": "system",
                    "content": (
                        "Create accurate college short-answer questions. "
                        "Return valid JSON only."
                    )
                },
                {
                    "role": "user",
                    "content": prompt
                }
            ],
            temperature=0.3
        )

        result = parse_ai_json(
            response.choices[0].message.content
        )

        questions = result.get("questions", [])

        if (
            not isinstance(questions, list)
            or len(questions) != question_count
        ):
            raise ValueError("Unexpected number of questions.")

        for question in questions:
            if not isinstance(question, dict):
                raise ValueError("Invalid question format.")

            if not all(
                isinstance(question.get(key), str)
                and question[key].strip()
                for key in ("question", "expected_answer")
            ):
                raise ValueError("Missing question information.")

        return jsonify({
            "questions": questions
        })

    except Exception:
        app.logger.exception(
            "Short-answer question generation failed"
        )

        return jsonify({
            "error": "Could not generate questions. Please try again."
        }), 500


# =========================================================
# SUBMIT SHORT ANSWERS, AI EVALUATION, AND HISTORY
# =========================================================

@app.route("/api/quiz/short-answers/submit", methods=["POST"])
def submit_short_answers_quiz():
    data = request.get_json(silent=True) or {}

    questions = data.get("questions")
    answers = data.get("answers")
    lecture_id = data.get("lecture_id")

    if (
        not isinstance(questions, list)
        or not isinstance(answers, list)
        or len(questions) != len(answers)
        or not 1 <= len(questions) <= 10
    ):
        return jsonify({
            "error": "Invalid quiz submission."
        }), 400

    try:
        lecture_id = int(lecture_id)
    except (TypeError, ValueError):
        return jsonify({
            "error": "Please select a valid lecture."
        }), 400

    if lecture_id <= 0:
        return jsonify({
            "error": "Please select a valid lecture."
        }), 400

    lecture = get_lecture_for_quiz(lecture_id)

    if lecture is None:
        return jsonify({
            "error": "Lecture not found."
        }), 404

    if client is None:
        return jsonify({
            "error": "AI is not configured. Check your .env file."
        }), 500

    validated_questions = []

    for question, answer in zip(questions, answers):
        if not isinstance(question, dict) or not isinstance(answer, str):
            return jsonify({
                "error": "Invalid answer data."
            }), 400

        question_text = question.get("question")
        expected_answer = question.get("expected_answer")

        if not all(
            isinstance(value, str) and value.strip()
            for value in (question_text, expected_answer)
        ):
            return jsonify({
                "error": "Invalid question data."
            }), 400

        validated_questions.append({
            "question": question_text,
            "expected_answer": expected_answer,
            "student_answer": answer.strip()[:2000]
        })

    grading_prompt = f"""
Evaluate the student's answers using ONLY the expected answers below.

Questions and answers:
{json.dumps(validated_questions, ensure_ascii=False)}

Return valid JSON only:
{{
  "results": [
    {{
      "score": 2,
      "max_score": 2,
      "feedback": "Correct. You identified the main function."
    }}
  ]
}}

Rules:
- Each question is worth 2 marks.
- Give 2 for correct, 1 for partially correct, and 0 for incorrect
  or empty answers.
- Accept different wording when the meaning is correct.
- Do not require exact wording.
- Do not award marks for unrelated answers.
- Keep feedback brief and constructive.
- Return exactly one result per question.
"""

    try:
        response = client.chat.completions.create(
            model="openai/gpt-oss-20b",
            messages=[
                {
                    "role": "system",
                    "content": (
                        "You are a fair college quiz evaluator. "
                        "Evaluate meaning rather than exact wording. "
                        "Return valid JSON only."
                    )
                },
                {
                    "role": "user",
                    "content": grading_prompt
                }
            ],
            temperature=0.1
        )

        result = parse_ai_json(
            response.choices[0].message.content
        )

        ai_results = result.get("results", [])

        if (
            not isinstance(ai_results, list)
            or len(ai_results) != len(questions)
        ):
            raise ValueError("Invalid grading results.")

        cleaned_results = []
        total_score = 0
        max_score = len(questions) * 2

        for index, item in enumerate(ai_results):
            if not isinstance(item, dict):
                raise ValueError("Invalid grading item.")

            score = item.get("score")
            feedback = item.get("feedback")

            if (
                isinstance(score, bool)
                or not isinstance(score, int)
                or score not in (0, 1, 2)
                or not isinstance(feedback, str)
            ):
                raise ValueError("Invalid score or feedback.")

            total_score += score

            cleaned_results.append({
                "question": validated_questions[index]["question"],
                "student_answer": validated_questions[index]["student_answer"],
                "expected_answer": validated_questions[index]["expected_answer"],
                "score": score,
                "max_score": 2,
                "feedback": feedback[:500]
            })

        save_quiz_attempt(
            "Short Answers",
            lecture_id,
            lecture["title"],
            total_score,
            max_score
        )

        percentage = round(
            (total_score / max_score) * 100,
            2
        )

        return jsonify({
            "score": total_score,
            "max_score": max_score,
            "percentage": percentage,
            "results": cleaned_results
        })

    except Exception:
        app.logger.exception("Short-answer grading failed")

        return jsonify({
            "error": "Could not evaluate answers. Please try again."
        }), 500


# =========================================================
# QUIZ HISTORY API
# =========================================================

@app.route("/api/quiz/history", methods=["GET"])
def get_quiz_history():
    connection = get_db_connection()

    try:
        rows = connection.execute("""
            SELECT
                id,
                quiz_type,
                lecture_id,
                lecture_title,
                score,
                max_score,
                percentage,
                attempted_at
            FROM quiz_history
            WHERE user_id = ?
            ORDER BY id DESC
        """, (get_current_user()["id"],)).fetchall()

        attempts = [dict(row) for row in rows]
        total_attempts = len(attempts)

        average_score = (
            round(
                sum(float(item["percentage"]) for item in attempts)
                / total_attempts,
                2
            )
            if total_attempts > 0
            else 0
        )

        return jsonify({
            "total_attempts": total_attempts,
            "average_score": average_score,
            "attempts": attempts
        })

    finally:
        connection.close()


# =========================================================
# QUIZ HISTORY DASHBOARD PAGE
# =========================================================

@app.route("/quiz/history")
def quiz_history_page():
    return render_template("quiz_history.html")


# =========================================================
# FILE TOO LARGE ERROR HANDLER
# =========================================================

@app.errorhandler(413)
def file_too_large(error):
    return (
        "The audio file is too large. "
        "Please choose a file smaller than 25 MB.",
        413
    )


# =========================================================
# INITIALIZE DATABASE AND START APPLICATION
# =========================================================

initialize_database()

if __name__ == "__main__":
    app.run(debug=True)